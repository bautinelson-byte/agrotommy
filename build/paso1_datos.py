"""Paso 1: arma Estructura EV consolidada, hoja Dólar BCRA y Gastos HA sobre el input 2026."""
import collections
import datetime as dt
import json

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

UP = '/root/.claude/uploads/5d49f2be-b7b1-55e1-b46a-c27cd044223b/'
INPUT = UP + 'a9e0d164-Adm._-_Estr._-_Gan._2026.xlsx'
PREV = UP + '446f283c-Gastos_Estructura_2024-2025.xlsx'
OUT = 'paso1.xlsx'

HA_DT, HA_EV = 840, 1605
F = 'Arial'
BOLD = Font(name=F, bold=True)
NORM = Font(name=F)
BLUE = Font(name=F, color='0000FF')
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
NEW_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
NUM = '#,##0;(#,##0);-'
NUM2 = '#,##0.00;(#,##0.00);-'
PCT = '0.0%;(0.0%);-'

# ---------- Mapeo de agrupamiento aprendido del ejercicio 24/25 ----------
prev = openpyxl.load_workbook(PREV)
prev_ev = [r for r in prev['Estructura EV'].iter_rows(min_row=2, values_only=True) if r[3]]
map_ctr = {}
map_art = collections.defaultdict(set)
for r in prev_ev:
    map_ctr[(r[0], r[3])] = r[2]
    map_art[r[3]].add(r[2])


def agrupar(ctr, art):
    """Devuelve (Cuenta agrupada, nueva?)."""
    if art in ("Electricidad", "Gas"):
        return art, True  # el cliente las va a clasificar
    if art in ('Electricidad', 'Gas'):
        return art, True  # las clasifica el cliente
    if (ctr, art) in map_ctr:
        return map_ctr[(ctr, art)], False
    return art, True  # sin clasificar: queda con su nombre para que lo revise el cliente


# ---------- Dólar ----------
tc = json.load(open('tc.json'))  # 'YYYY-MM' -> [fecha, valor]

wb = openpyxl.load_workbook(INPUT)
FUENTES = [  # hoja, Ctro. De Costo, Campo
    ('Estructura EV', 'Estructura', 'El Venado'),
    ('Estructura DT', 'Estructura', 'Don Tomas'),
    ('Ganaderia', 'Ganaderia', 'El Venado'),
    ('Administracion', 'Administración', 'Administración'),
]
rows = []
for hoja, ctr, campo in FUENTES:
    for r in wb[hoja].iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in r):
            continue
        art = (r[0] or '').strip()
        cuenta, nueva = agrupar(ctr, art)
        rows.append([ctr, campo, cuenta, art, r[1], r[2], r[3], r[4], r[5],
                     (r[6] or '').strip(), r[7] or 0, r[8] or 0, nueva])

# ---------- Hoja Dólar BCRA ----------
wd = wb.create_sheet('Dólar BCRA')
wd['A1'] = 'Tipo de cambio BCRA Com. A 3500 (mayorista) - último día hábil de cada mes'
wd['A1'].font = BOLD
wd['A2'] = ('Fuente: API BCRA estadisticascambiarias v1.0 (api.bcra.gob.ar), mismo criterio que el ejercicio 24/25. '
            'Para un mes nuevo, agregar una fila.')
wd['A2'].font = Font(name=F, italic=True, color='808080')
for c, h in enumerate(['Período', 'Fecha del dato', 'ARS por USD'], 1):
    x = wd.cell(row=4, column=c, value=h)
    x.font, x.fill = BOLD, HEAD_FILL
for i, k in enumerate(sorted(tc), 5):
    y, m = map(int, k.split('-'))
    wd.cell(row=i, column=1, value=dt.date(y, m, 1)).number_format = 'mmm-yyyy'
    wd.cell(row=i, column=2, value=dt.date.fromisoformat(tc[k][0])).number_format = 'dd/mm/yyyy'
    v = wd.cell(row=i, column=3, value=tc[k][1])
    v.number_format, v.font = '#,##0.00', BLUE
    wd.cell(row=i, column=1).font = wd.cell(row=i, column=2).font = NORM
last_tc = 4 + len(tc)
for col, w in zip('ABC', (12, 15, 14)):
    wd.column_dimensions[col].width = w

# ---------- Estructura EV consolidada ----------
del wb['Estructura EV']
ws = wb.create_sheet('Estructura EV', 0)
HEAD = ['Ctro. De Costo', 'Campo', 'Cuenta', 'Articulo', 'Fecha', 'Documento', 'Numero', 'Descripcion',
        'Tercero', 'Centro de Costo', 'Debe', 'Haber', 'Importe Bruto', 'Año', 'Mes', 'Dólar BCRA',
        'Importe USD', 'Revisar']
for c, h in enumerate(HEAD, 1):
    x = ws.cell(row=1, column=c, value=h)
    x.font, x.fill = BOLD, HEAD_FILL
TCR = f"'Dólar BCRA'!$A$5:$A${last_tc}"
TCV = f"'Dólar BCRA'!$C$5:$C${last_tc}"
for i, r in enumerate(rows, 2):
    for c, v in enumerate(r[:12], 1):
        ws.cell(row=i, column=c, value=v).font = NORM
    ws.cell(row=i, column=5).number_format = 'dd/mm/yyyy'
    ws.cell(row=i, column=11).number_format = NUM2
    ws.cell(row=i, column=12).number_format = NUM2
    ws.cell(row=i, column=13, value=f'=K{i}-L{i}').number_format = NUM2
    ws.cell(row=i, column=14, value=f'=YEAR(E{i})')
    ws.cell(row=i, column=15, value=f'=TEXT(MONTH(E{i}),"00")&"-"&CHOOSE(MONTH(E{i}),"ene","feb","mar","abr","may","jun","jul","ago","sep","oct","nov","dic")')
    ws.cell(row=i, column=16, value=f'=INDEX({TCV},MATCH(DATE(YEAR(E{i}),MONTH(E{i}),1),{TCR},0))').number_format = '#,##0.00'
    ws.cell(row=i, column=17, value=f'=M{i}/P{i}').number_format = NUM2
    if r[12]:
        ws.cell(row=i, column=18, value='Cuenta nueva: clasificar').fill = NEW_FILL
        ws.cell(row=i, column=3).fill = NEW_FILL
    for c in range(13, 19):
        ws.cell(row=i, column=c).font = NORM
NR = len(rows) + 1
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:R{NR}'
for c, w in enumerate([14, 14, 34, 34, 11, 12, 9, 34, 30, 22, 15, 15, 15, 7, 8, 11, 13, 22], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws['C1'].comment = Comment('Agrupamiento del ejercicio 24/25. Las cuentas que no existían ese año quedan con su '
                           'nombre original y marcadas en amarillo (columna Revisar) para que el cliente las clasifique.',
                           'Claude')

# ---------- Gastos HA ----------
EV = "'Estructura EV'"
rng = lambda col: f'{EV}!${col}$2:${col}${NR}'
A_CTR, A_CAMPO, A_CTA, A_ART, A_IMP, A_USD = rng('A'), rng('B'), rng('C'), rng('D'), rng('M'), rng('Q')

prev_ha = prev['Gastos HA']
prev_pes = {prev_ha.cell(r, 1).value: prev_ha.cell(r, 5).value for r in range(4, 23)}
prev_usd = {prev_ha.cell(r, 10).value: (prev_ha.cell(r, 11).value or 0) + (prev_ha.cell(r, 12).value or 0)
            + (prev_ha.cell(r, 13).value or 0) for r in range(4, 23)}
# los totales del pesos 24/25 en el archivo viejo son fórmulas: recalcular desde las columnas B:D
prev_pes = {prev_ha.cell(r, 1).value: sum(prev_ha.cell(r, c).value or 0 for c in (2, 3, 4)) for r in range(4, 23)}
hist = {}  # campo -> {cuenta: [21/22,22/23,23/24,24/25]}
for titulo, r0 in (('Administración', 38), ('Don Tomás', 61), ('El Venado', 85)):
    hist[titulo] = {prev_ha.cell(r, 18).value: [prev_ha.cell(r, c).value for c in range(19, 23)]
                    for r in range(r0, r0 + 19)}

cuentas = sorted({r[2] for r in rows if r[0] != 'Ganaderia'} | set(prev_pes), key=str.lower)
gh = wb.create_sheet('Gastos HA', 0)


def head(row, col, vals):
    for j, v in enumerate(vals):
        x = gh.cell(row=row, column=col + j, value=v)
        x.font, x.fill = BOLD, HEAD_FILL
        x.alignment = Alignment(horizontal='center', wrap_text=True)


gh['A1'] = 'Gastos de Estructura y Administración - Ejercicio 2025-2026 (jul-25 a jun-26)'
gh['A1'].font = Font(name=F, bold=True, size=12)
gh['A2'] = 'Hectáreas'
gh['B2'], gh['C2'] = 'Don Tomás', HA_DT
gh['D2'], gh['E2'] = 'El Venado', HA_EV
gh['F2'], gh['G2'] = 'Total', '=C2+E2'
for c in ('A2', 'B2', 'D2', 'F2', 'G2'):
    gh[c].font = BOLD
gh['C2'].font = gh['E2'].font = BLUE
gh['C2'].comment = Comment('Hectáreas informadas por el usuario (mismas que 24/25).', 'Claude')

# Tabla en pesos (A) y en dólares (J)
r0 = 4
gh.cell(row=r0, column=1, value='En pesos').font = BOLD
gh.cell(row=r0, column=10, value='En dólares (BCRA fin de mes)').font = BOLD
head(r0 + 1, 1, ['Cuenta', 'Administración', 'Don Tomas', 'El Venado', 'Total 25/26', '2024-2025', 'Var. %'])
head(r0 + 1, 10, ['Cuenta', 'Administración', 'Don Tomas', 'El Venado', 'Total 25/26', '2024-2025', 'Var. %'])
first = r0 + 2
for k, cta in enumerate(cuentas):
    i = first + k
    for base, rngv, prevd in ((1, A_IMP, prev_pes), (10, A_USD, prev_usd)):
        L = lambda c: get_column_letter(base + c)
        gh.cell(row=i, column=base, value=cta).font = NORM
        for j, campo in enumerate(('Administración', 'Don Tomas', 'El Venado'), 1):
            gh.cell(row=i, column=base + j, value=(
                f'=SUMIFS({rngv},{A_CTA},{L(0)}{i},{A_CAMPO},"{campo}",{A_CTR},"<>Ganaderia")'))
        gh.cell(row=i, column=base + 4, value=f'=SUM({L(1)}{i}:{L(3)}{i})')
        p = gh.cell(row=i, column=base + 5, value=prevd.get(cta))
        p.font = BLUE
        gh.cell(row=i, column=base + 6, value=f'=IF(N({L(5)}{i})=0,"",({L(4)}{i}-{L(5)}{i})/{L(5)}{i})')
        if cta not in prevd:
            gh.cell(row=i, column=base).fill = NEW_FILL
last = first + len(cuentas) - 1
tot = last + 1
for base in (1, 10):
    L = lambda c: get_column_letter(base + c)
    gh.cell(row=tot, column=base, value='Total general').font = BOLD
    for j in range(1, 6):
        gh.cell(row=tot, column=base + j, value=f'=SUM({L(j)}{first}:{L(j)}{last})').font = BOLD
    gh.cell(row=tot, column=base + 6, value=f'=IF({L(5)}{tot}=0,"",({L(4)}{tot}-{L(5)}{tot})/{L(5)}{tot})').font = BOLD
for i in range(first, tot + 1):
    for c in list(range(2, 7)) + list(range(11, 16)):
        gh.cell(row=i, column=c).number_format = NUM
    gh.cell(row=i, column=7).number_format = PCT
    gh.cell(row=i, column=16).number_format = PCT
    for c in range(1, 17):
        cell = gh.cell(row=i, column=c)
        if cell.font != BLUE and not cell.font.bold:
            cell.font = NORM
gh.cell(row=tot + 1, column=1, value='Columna 2024-2025: valores del archivo "Gastos Estructura 2024-2025" (Gastos HA). '
        'En amarillo: cuentas sin dato el año anterior.').font = Font(name=F, italic=True, color='808080')

# Por hectárea (USD/ha) por campo, con histórico
r = tot + 4
gh.cell(row=r, column=1, value='Dólares por hectárea').font = Font(name=F, bold=True, size=12)
gh.cell(row=r + 1, column=1, value='Administración se reparte sobre el total de hectáreas.').font = Font(name=F, italic=True, color='808080')
r += 3
bloques = []
for titulo, col_usd, ha in (('Administración', 'K', '$G$2'), ('Don Tomás', 'L', '$C$2'), ('El Venado', 'M', '$E$2')):
    gh.cell(row=r, column=1, value=titulo).font = BOLD
    gh.cell(row=r, column=8, value=f'{titulo} / ha').font = BOLD
    head(r + 1, 1, ['Cuenta', '21/22', '22/23', '23/24', '24/25', '25/26'])
    head(r + 1, 8, ['21/22', '22/23', '23/24', '24/25', '25/26'])
    b0 = r + 2
    for k, cta in enumerate(cuentas):
        i = b0 + k
        gh.cell(row=i, column=1, value=cta).font = NORM
        h = hist[titulo].get(cta, [None] * 4)
        for j in range(4):
            x = gh.cell(row=i, column=2 + j, value=h[j])
            x.font, x.number_format = BLUE, NUM
        x = gh.cell(row=i, column=6, value=f'={col_usd}{first + k}')
        x.number_format, x.font = NUM, Font(name=F, color='008000')
        for j in range(5):
            y = gh.cell(row=i, column=8 + j, value=f'=IF({get_column_letter(2 + j)}{i}="","",{get_column_letter(2 + j)}{i}/{ha})')
            y.number_format, y.font = NUM2, NORM
    bt = b0 + len(cuentas)
    gh.cell(row=bt, column=1, value='Total general').font = BOLD
    for c in list(range(2, 7)) + list(range(8, 13)):
        L = get_column_letter(c)
        x = gh.cell(row=bt, column=c, value=f'=SUM({L}{b0}:{L}{bt - 1})')
        x.font, x.number_format = BOLD, (NUM if c < 8 else NUM2)
    bloques.append((titulo, bt))
    r = bt + 3

# Resumen USD/ha + inmobiliario y red vial
gh.cell(row=r, column=1, value='Resumen USD / ha').font = BOLD
head(r + 1, 1, ['', '21/22', '22/23', '23/24', '24/25', '25/26'])
nombres = {'Administración': 'Administración', 'Don Tomás': 'Estructura DT', 'El Venado': 'Estructura EV'}
for k, (titulo, bt) in enumerate(bloques):
    i = r + 2 + k
    gh.cell(row=i, column=1, value=nombres[titulo]).font = NORM
    for j in range(5):
        x = gh.cell(row=i, column=2 + j, value=f'={get_column_letter(8 + j)}{bt}')
        x.number_format, x.font = NUM2, NORM
r = r + 2 + len(bloques) + 1
prev_w = {prev_ha.cell(rr, 23).value: [prev_ha.cell(rr, c).value for c in range(24, 28)] for rr in (112, 113, 115, 116)}
for k, (nombre, arts, campo, ha) in enumerate((
        ('Inmobiliario DT', ['Impuesto Inmobiliario'], 'Don Tomas', '$C$2'),
        ('Red Vial DT', ['Red Vial Teodelina', 'Red Vial Ameghino'], 'Don Tomas', '$C$2'),
        ('Inmobiliario EV', ['Impuesto Inmobiliario'], 'El Venado', '$E$2'),
        ('Red Vial EV', ['Red Vial Teodelina', 'Red Vial Ameghino'], 'El Venado', '$E$2'))):
    i = r + k
    gh.cell(row=i, column=1, value=nombre).font = NORM
    for j, v in enumerate(prev_w[nombre]):
        x = gh.cell(row=i, column=2 + j, value=v)
        x.font, x.number_format = BLUE, NUM2
    f = '+'.join(f'SUMIFS({A_USD},{A_ART},"{a}",{A_CAMPO},"{campo}",{A_CTR},"Estructura")' for a in arts)
    x = gh.cell(row=i, column=6, value=f'=({f})/{ha}')
    x.font, x.number_format = NORM, NUM2
gh.cell(row=r + 5, column=1, value='Históricos (21/22 a 24/25) copiados del archivo 2024-2025.').font = Font(name=F, italic=True, color='808080')
gh.column_dimensions['A'].width = 42
gh.column_dimensions['J'].width = 42
for c in 'BCDEFGHIKLMNOP':
    gh.column_dimensions[c].width = 15
gh.freeze_panes = 'B4'

# ---------- Comparación de Sueldos contra 24/25 (en Gastos HA, al final) ----------
r += 8
gh.cell(row=r, column=1, value='Sueldos (Estructura) por mes: 25/26 vs 24/25 en pesos').font = Font(name=F, bold=True, size=12)
meses = [(2025, m) for m in range(7, 13)] + [(2026, m) for m in range(1, 7)]
nom = ['jul', 'ago', 'sep', 'oct', 'nov', 'dic', 'ene', 'feb', 'mar', 'abr', 'may', 'jun']
prev_su = collections.defaultdict(float)
for x in prev_ev:
    if x[0] == 'Estructura' and x[2] == 'Sueldos' and hasattr(x[4], 'year'):
        v = x[10]
        if isinstance(v, str):
            v = eval(v.lstrip('='))
        prev_su[(x[3], x[4].year + 1, x[4].month)] += v or 0
arts_su = sorted({x[3] for x in rows if x[0] == 'Estructura' and x[2] == 'Sueldos'} | {k[0] for k in prev_su})
A_FECHA = rng('E')
for bloque, titulo in enumerate(('25/26', '24/25', 'Variación %')):
    rb = r + 2 + bloque * (len(arts_su) + 4)
    gh.cell(row=rb, column=1, value=titulo).font = BOLD
    head(rb + 1, 1, ['Articulo'] + nom + ['Total'])
    for k, a in enumerate(arts_su + ['Total']):
        i = rb + 2 + k
        gh.cell(row=i, column=1, value=a).font = BOLD if a == 'Total' else NORM
        for j, (y, m) in enumerate(meses):
            L = get_column_letter(2 + j)
            if a == 'Total' and bloque < 2:
                v = f'=SUM({L}{rb + 2}:{L}{i - 1})'
            elif bloque == 0:
                d1, d2 = dt.date(y, m, 1), dt.date(y + (m == 12), m % 12 + 1, 1)
                v = (f'=SUMIFS({A_IMP},{A_ART},$A{i},{A_CTA},"Sueldos",{A_CTR},"Estructura",'
                     f'{A_FECHA},">="&DATE({d1.year},{d1.month},1),{A_FECHA},"<"&DATE({d2.year},{d2.month},1))')
            elif bloque == 1:
                v = round(prev_su.get((a, y, m), 0), 2)
            else:
                cur, pre = i - 2 * (len(arts_su) + 4), i - (len(arts_su) + 4)
                v = f'=IF(N({L}{pre})=0,"",({L}{cur}-{L}{pre})/{L}{pre})'
            x = gh.cell(row=i, column=2 + j, value=v)
            x.number_format = PCT if bloque == 2 else NUM
            x.font = BLUE if bloque == 1 and a != 'Total' else NORM
        T = get_column_letter(14)
        if bloque < 2:
            gh.cell(row=i, column=14, value=f'=SUM(B{i}:M{i})').number_format = NUM
        else:
            cur, pre = i - 2 * (len(arts_su) + 4), i - (len(arts_su) + 4)
            gh.cell(row=i, column=14, value=f'=IF(N(N{pre})=0,"",(N{cur}-N{pre})/N{pre})').number_format = PCT

# ---------- Orden de hojas ----------
tr = wb.create_sheet('Tabla Resumen', 1)
order = ['Gastos HA', 'Tabla Resumen', 'Estructura EV', 'Estructura DT', 'Ganaderia', 'Administracion', 'Dólar BCRA']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
print('filas EV', len(rows), 'nuevas', sum(r[12] for r in rows), 'cuentas', len(cuentas), 'NR', NR)
print('sin clasificar:', sorted({(r[0], r[3]) for r in rows if r[12]}))
